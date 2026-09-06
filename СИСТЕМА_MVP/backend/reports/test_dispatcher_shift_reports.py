from datetime import datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.production_time import production_shift_context
from downtimes.models import DowntimeEvent, DowntimeReason
from references.models import DumpPoint, Equipment, EquipmentModel, EquipmentType, RockType
from shifts.models import EmployeeShift
from trips.models import DispatcherActionLog, Trip, TripStatus
from users.models import Employee, EmployeeAccess, Role

from .dispatcher_shift_forms import build_dispatcher_shift_report


class DispatcherShiftReportTests(TestCase):
    def setUp(self):
        role = Role.objects.create(code='dispatcher', name='Диспетчер')
        self.dispatcher = Employee.objects.create(
            full_name='Диспетчер отчётов',
            status=Employee.Status.ACTIVE,
        )
        self.access = EmployeeAccess.objects.create(
            employee=self.dispatcher,
            role=role,
            access_code='991100',
            status=EmployeeAccess.Status.ACTIVATED,
            is_active=True,
        )
        session = self.client.session
        session['employee_access_id'] = self.access.id
        session.save()

        truck_type = EquipmentType.objects.create(name='Самосвал')
        excavator_type = EquipmentType.objects.create(name='Экскаватор')
        belaz_model = EquipmentModel.objects.create(equipment_type=truck_type, name='БелАЗ 75131')
        nhl_model = EquipmentModel.objects.create(equipment_type=truck_type, name='NHL TR100')
        excavator_model = EquipmentModel.objects.create(equipment_type=excavator_type, name='Sany 8800')
        self.belaz = Equipment.objects.create(equipment_type=truck_type, model=belaz_model, garage_number='10')
        self.nhl = Equipment.objects.create(equipment_type=truck_type, model=nhl_model, garage_number='54')
        self.idle_truck = Equipment.objects.create(equipment_type=truck_type, model=belaz_model, garage_number='14')
        self.excavator = Equipment.objects.create(equipment_type=excavator_type, model=excavator_model, garage_number='1')
        self.rock = RockType.objects.create(
            name='Первичная сульфидная',
            density=Decimal('2.5000'),
        )
        self.dump_point = DumpPoint.objects.create(name='ККД')
        self.selected_date = production_shift_context().production_date
        self.shift_at = timezone.make_aware(
            datetime.combine(self.selected_date, time(10, 0)),
            timezone.get_current_timezone(),
        )
        operator = Employee.objects.create(full_name='Машинист')
        self.loading_shift = EmployeeShift.objects.create(
            employee=operator,
            shift_type='day',
            equipment=self.excavator,
            opened_at=self.shift_at,
        )
        self.trip = self.create_trip(self.belaz, '50.00', '3.00', completed_at=self.shift_at)
        self.create_trip(self.belaz, '50.00', '5.00', completed_at=self.shift_at + timedelta(minutes=30))
        self.create_trip(self.nhl, '40.00', '2.00', completed_at=self.shift_at + timedelta(hours=1))

        reason = DowntimeReason.objects.create(name='Монтаж рамы', equipment_type=truck_type, is_critical=True)
        DowntimeEvent.objects.create(
            equipment=self.idle_truck,
            reason=reason,
            started_at=self.shift_at,
            ended_at=self.shift_at + timedelta(hours=2),
            comment='сварка рамы',
        )

    def create_trip(self, truck, volume, distance, *, completed_at):
        trip = Trip.objects.create(
            excavator=self.excavator,
            truck=truck,
            loading_shift=self.loading_shift,
            rock_type=self.rock,
            dump_point=self.dump_point,
            planned_volume_m3=Decimal('10000.00'),
            volume_m3=Decimal(volume),
            transport_distance_km=Decimal(distance),
            loading_horizon='90',
            loading_block='192',
            status=TripStatus.COMPLETED,
            completed_at=completed_at,
        )
        Trip.objects.filter(pk=trip.pk).update(created_at=completed_at)
        trip.created_at = completed_at
        return trip

    def test_report_calculates_weighted_distance_and_m3km(self):
        report = build_dispatcher_shift_report(self.selected_date, 'day')
        belaz_row = next(row for row in report['truck_rows'] if row['equipment'].id == self.belaz.id)
        idle_row = next(row for row in report['truck_rows'] if row['equipment'].id == self.idle_truck.id)

        self.assertEqual(belaz_row['trip_count'], 2)
        self.assertEqual(belaz_row['volume'], Decimal('100.00'))
        self.assertEqual(belaz_row['m3km'], Decimal('400.0000'))
        self.assertEqual(belaz_row['distance'], Decimal('4.0'))
        self.assertEqual(report['truck_totals']['trip_count'], 3)
        self.assertEqual(report['truck_totals']['volume'], Decimal('140.00'))
        self.assertEqual(report['truck_totals']['m3km'], Decimal('480.0000'))
        self.assertEqual(report['truck_totals']['distance'], Decimal('3.4'))
        self.assertEqual(idle_row['trip_count'], 0)
        self.assertIn('Монтаж рамы—2ч', idle_row['notes'])
        self.assertEqual(next(row for row in report['truck_rows'] if row['equipment'].id == self.nhl.id)['fleet'], 'nhl')

    def test_report_builds_excavation_customer_rows(self):
        report = build_dispatcher_shift_report(self.selected_date, 'day')

        self.assertEqual(len(report['excavation_rows']), 3)
        self.assertEqual(report['excavation_totals']['volume'], Decimal('140.00'))
        self.assertEqual(report['excavation_totals']['trip_count'], 3)
        self.assertTrue(all(row['rock_type'] == 'Первичная сульфидная' for row in report['excavation_rows']))
        self.assertTrue(all(row['horizon'] == '90' for row in report['excavation_rows']))

    def test_hourly_report_uses_loading_hour_and_only_actual_hour_dump_points(self):
        second_point = DumpPoint.objects.create(name='СКДР')
        nhl_trip = Trip.objects.filter(truck=self.nhl).get()
        nhl_trip.actual_dump_point = second_point
        nhl_trip.save(update_fields=['actual_dump_point'])
        in_transit = Trip.objects.create(
            excavator=self.excavator,
            truck=self.belaz,
            loading_shift=self.loading_shift,
            rock_type=self.rock,
            dump_point=second_point,
            volume_m3=Decimal('25.00'),
            status=TripStatus.LOADED_WAITING_UNLOAD,
        )
        Trip.objects.filter(pk=in_transit.pk).update(
            created_at=self.shift_at + timedelta(hours=2, minutes=15),
        )

        report = build_dispatcher_shift_report(self.selected_date, 'day')
        hourly = report['hourly']
        first_hour = next(hour for hour in hourly['hours'] if hour['label'] == '11.00')
        second_hour = next(hour for hour in hourly['hours'] if hour['label'] == '12.00')
        third_hour = next(hour for hour in hourly['hours'] if hour['label'] == '13.00')
        excavator_row = next(row for row in hourly['rows'] if row['equipment'].id == self.excavator.id)

        self.assertEqual([point['label'] for point in first_hour['points']], ['ККД'])
        self.assertEqual([point['label'] for point in second_hour['points']], ['СКДР'])
        self.assertNotIn('СКДР', [point['label'] for point in first_hour['points']])
        self.assertNotIn('ККД', [point['label'] for point in second_hour['points']])
        self.assertEqual(first_hour['trip_count'], 2)
        self.assertEqual(second_hour['trip_count'], 1)
        self.assertEqual(third_hour['trip_count'], 1)
        self.assertEqual(excavator_row['trip_count'], 4)
        self.assertEqual(hourly['volume'], Decimal('165.00'))

    def test_excavation_downtimes_are_grouped_by_reason_and_zero_totals_hidden(self):
        waiting, _ = DowntimeReason.objects.get_or_create(
            name='Ожидание самосвалов',
            defaults={'equipment_type': self.excavator.equipment_type},
        )
        inspection, _ = DowntimeReason.objects.get_or_create(
            name='Тестовый осмотр для отчёта',
            defaults={'equipment_type': self.excavator.equipment_type},
        )
        DowntimeEvent.objects.create(
            equipment=self.excavator,
            reason=waiting,
            started_at=self.shift_at,
            ended_at=self.shift_at + timedelta(seconds=30),
            comment='Автоматически по производственному событию',
        )
        DowntimeEvent.objects.create(
            equipment=self.excavator,
            reason=waiting,
            started_at=self.shift_at + timedelta(minutes=1),
            ended_at=self.shift_at + timedelta(minutes=1, seconds=45),
            comment='Автоматически по производственному событию',
        )
        DowntimeEvent.objects.create(
            equipment=self.excavator,
            reason=inspection,
            started_at=self.shift_at + timedelta(minutes=2),
            ended_at=self.shift_at + timedelta(minutes=2, seconds=20),
        )

        report = build_dispatcher_shift_report(self.selected_date, 'day')

        self.assertTrue(report['excavation_rows'])
        downtime_values = [row['downtime'] for row in report['excavation_rows']]
        self.assertEqual(downtime_values.count('Ожидание самосвалов—1 мин'), 1)
        self.assertTrue(all(not value for value in downtime_values[1:]))
        self.assertNotIn('0 мин', '; '.join(downtime_values))
        self.assertNotIn('Автоматически', '; '.join(downtime_values))
        self.assertNotIn('Тестовый осмотр', '; '.join(downtime_values))

    def test_truck_downtimes_are_grouped_by_reason_and_zero_totals_hidden(self):
        waiting, _ = DowntimeReason.objects.get_or_create(
            name='Ожидание разгрузки',
            defaults={'equipment_type': self.belaz.equipment_type},
        )
        inspection, _ = DowntimeReason.objects.get_or_create(
            name='Тестовый осмотр самосвала для отчёта',
            defaults={'equipment_type': self.belaz.equipment_type},
        )
        DowntimeEvent.objects.create(
            equipment=self.belaz,
            reason=waiting,
            started_at=self.shift_at,
            ended_at=self.shift_at + timedelta(seconds=30),
            comment='Автоматически по производственному событию',
        )
        DowntimeEvent.objects.create(
            equipment=self.belaz,
            reason=waiting,
            started_at=self.shift_at + timedelta(minutes=1),
            ended_at=self.shift_at + timedelta(minutes=1, seconds=45),
            comment='Автоматически по производственному событию',
        )
        DowntimeEvent.objects.create(
            equipment=self.belaz,
            reason=inspection,
            started_at=self.shift_at + timedelta(minutes=2),
            ended_at=self.shift_at + timedelta(minutes=2, seconds=20),
        )

        report = build_dispatcher_shift_report(self.selected_date, 'day')
        truck_row = next(
            row for row in report['truck_rows']
            if row['equipment'].id == self.belaz.id
        )

        self.assertEqual(truck_row['notes'], 'Ожидание разгрузки—1 мин')
        self.assertEqual(truck_row['notes'].count('Ожидание разгрузки'), 1)
        self.assertNotIn('0 мин', truck_row['notes'])
        self.assertNotIn('Автоматически', truck_row['notes'])
        self.assertNotIn('Тестовый осмотр', truck_row['notes'])

    def test_reports_hub_and_both_forms_are_available(self):
        params = {'date': self.selected_date.isoformat(), 'shift_type': 'day'}
        hub = self.client.get(reverse('dispatcher_reports'), params)
        trucks = self.client.get(reverse('dispatcher_shift_trucks'), params)
        excavation = self.client.get(reverse('dispatcher_shift_excavation'), params)
        hourly = self.client.get(reverse('dispatcher_shift_hourly'), params)

        self.assertEqual(hub.status_code, 200)
        self.assertContains(hub, 'Итоги смены по самосвалам')
        self.assertContains(hub, 'Работа выемочного оборудования')
        self.assertContains(hub, 'Почасовая сводка смены')
        self.assertEqual(trucks.status_code, 200)
        self.assertContains(trucks, 'dispatcher-shift-report-screen')
        self.assertContains(trucks, 'data-dispatcher-shift-report-live')
        self.assertContains(trucks, 'window.applyOperationalStateRefresh')
        self.assertContains(trucks, 'refreshDispatcherShiftReport')
        self.assertContains(trucks, 'name: "dispatcher-shift-report", role: "dispatcher", mode: "custom"')
        self.assertContains(trucks, 'this.form.requestSubmit()', count=2)
        self.assertNotContains(trucks, '>Показать</button>')
        self.assertContains(trucks, 'м³×км')
        self.assertContains(trucks, 'Корректировать')
        self.assertEqual(excavation.status_code, 200)
        self.assertContains(excavation, 'dispatcher-shift-report-screen')
        self.assertContains(excavation, 'data-dispatcher-shift-report-live')
        self.assertContains(excavation, 'Тип грунта')
        self.assertContains(excavation, 'Место разгрузки')
        self.assertEqual(hourly.status_code, 200)
        self.assertContains(hourly, 'Почасовая сводка смены')
        self.assertContains(hourly, 'ККД')
        self.assertContains(hourly, 'data-dispatcher-shift-report-live')

    def test_both_shift_reports_expose_live_operational_fragment(self):
        params = {
            'date': self.selected_date.isoformat(),
            'shift_type': 'day',
            '_operational_fragment': 'dispatcher-shift-report',
            '_operational_version': '0',
        }

        for route_name in ('dispatcher_shift_trucks', 'dispatcher_shift_excavation', 'dispatcher_shift_hourly'):
            with self.subTest(route_name=route_name):
                response = self.client.get(
                    reverse(route_name),
                    params,
                    HTTP_ACCEPT='application/json',
                    HTTP_X_REQUESTED_WITH='XMLHttpRequest',
                )

                self.assertEqual(response.status_code, 200)
                self.assertEqual(response['Cache-Control'], 'no-store')
                self.assertEqual(response['X-Operational-Fragment'], 'operational-fragment-v1')
                payload = response.json()
                self.assertEqual(payload['contract'], 'operational-fragment-v1')
                self.assertEqual(payload['screen'], 'dispatcher-shift-report')
                self.assertIsInstance(payload['version'], int)
                self.assertIn('data-dispatcher-shift-report-live', payload['html'])
                self.assertNotIn('<script', payload['html'].lower())
                self.assertNotIn('<main', payload['html'].lower())

    def test_excel_exports_keep_numeric_values_and_familiar_sheets(self):
        params = {'date': self.selected_date.isoformat(), 'shift_type': 'day'}
        trucks_response = self.client.get(reverse('dispatcher_shift_trucks_export'), params)
        excavation_response = self.client.get(reverse('dispatcher_shift_excavation_export'), params)
        hourly_response = self.client.get(reverse('dispatcher_shift_hourly_export'), params)

        self.assertEqual(trucks_response.status_code, 200)
        truck_book = load_workbook(BytesIO(trucks_response.content), data_only=False)
        self.assertEqual(truck_book.sheetnames, ['Самосвалы'])
        truck_sheet = truck_book['Самосвалы']
        self.assertEqual(truck_sheet['A3'].value, '№ А/С')
        total_row = truck_sheet.max_row
        self.assertEqual(truck_sheet.cell(total_row, 2).value, 3)
        self.assertEqual(truck_sheet.cell(total_row, 4).value, 140)
        self.assertEqual(truck_sheet.cell(total_row, 5).value, 480)
        self.assertIsInstance(truck_sheet.cell(total_row, 4).value, (int, float))

        self.assertEqual(excavation_response.status_code, 200)
        excavation_book = load_workbook(BytesIO(excavation_response.content), data_only=False)
        self.assertEqual(excavation_book.sheetnames, ['Выемочное оборудование'])
        self.assertEqual(excavation_book.active['A3'].value, 'Тип грунта')

        self.assertEqual(hourly_response.status_code, 200)
        hourly_book = load_workbook(BytesIO(hourly_response.content), data_only=False)
        self.assertEqual(hourly_book.sheetnames, ['Почасовая сводка'])
        hourly_sheet = hourly_book.active
        self.assertEqual(hourly_sheet['A1'].value, 'Почасовая сводка смены')
        self.assertEqual(hourly_sheet['C3'].value, '08.00')
        self.assertIn('ККД', [cell.value for cell in hourly_sheet[4]])
        self.assertTrue(any(cell.data_type == 'f' for row in hourly_sheet.iter_rows() for cell in row))

    def test_dispatcher_correction_updates_source_and_writes_audit_log(self):
        response = self.client.post(reverse('dispatcher_shift_trucks'), {
            'date': self.selected_date.isoformat(),
            'shift_type': 'day',
            'trip_id': self.trip.id,
            'volume_m3': '55,5',
            'transport_distance_km': '3,2',
            'rock_type_id': self.rock.id,
            'actual_dump_point_id': self.dump_point.id,
            'loading_horizon': '95',
            'loading_block': '193',
            'downtime_text': 'ОП 20 мин',
            'note': 'Уточнено диспетчером',
            'correction_reason': 'Сверка со сводкой',
        })

        self.assertEqual(response.status_code, 302)
        self.trip.refresh_from_db()
        self.assertEqual(self.trip.volume_m3, Decimal('55.50'))
        self.assertEqual(self.trip.tonnage, Decimal('138.75'))
        self.assertEqual(self.trip.transport_distance_km, Decimal('3.20'))
        self.assertEqual(self.trip.loading_horizon, '95')
        action = DispatcherActionLog.objects.get(trip=self.trip, action_type='report_source_correction')
        self.assertEqual(action.actor, self.dispatcher)
        self.assertEqual(action.reason, 'Сверка со сводкой')
        journal = self.client.get(reverse('dispatcher_shift_log'), {'date': self.selected_date.isoformat()})
        self.assertContains(journal, 'Корректировка исходных данных отчёта')

    def test_correction_rejects_trip_from_another_shift(self):
        response = self.client.post(reverse('dispatcher_shift_trucks'), {
            'date': self.selected_date.isoformat(),
            'shift_type': 'night',
            'trip_id': self.trip.id,
            'volume_m3': '99',
            'transport_distance_km': '3',
            'rock_type_id': self.rock.id,
            'actual_dump_point_id': '',
            'loading_horizon': '90',
            'loading_block': '192',
            'downtime_text': '',
            'note': '',
            'correction_reason': 'Ошибочная смена',
        })

        self.assertEqual(response.status_code, 200)
        self.trip.refresh_from_db()
        self.assertEqual(self.trip.volume_m3, Decimal('50.00'))
        self.assertFalse(DispatcherActionLog.objects.filter(trip=self.trip, action_type='report_source_correction').exists())
