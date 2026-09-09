import re
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal, ROUND_HALF_UP

from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.production_time import (
    assigned_shift_open_bounds,
    production_shift_bounds,
)
from downtimes.models import DowntimeEvent
from references.models import Equipment
from shifts.models import EmployeeShift, ShiftType
from trips.models import OPEN_TRIP_STATUSES, Trip, TripStatus


ZERO = Decimal('0')
ONE_DECIMAL = Decimal('0.1')
THIN_BLACK = Side(style='thin', color='000000')
MEDIUM_BLACK = Side(style='medium', color='000000')


def natural_key(value):
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r'(\d+)', str(value))]


def decimal_sum(values):
    return sum((value or ZERO for value in values), ZERO)


def equipment_label(equipment):
    model_name = equipment.model.name if equipment.model_id else ''
    combined = f'{model_name} {equipment.garage_number}'.upper()
    garage_number = equipment.garage_number.strip()
    if 'NHL' in combined:
        return garage_number if 'NHL' in garage_number.upper() else f'NHL №{garage_number}'
    return garage_number if garage_number.startswith('№') else f'№{garage_number}'


def is_nhl(equipment):
    model_name = equipment.model.name if equipment.model_id else ''
    return 'NHL' in f'{model_name} {equipment.garage_number}'.upper()


def shift_meta(selected_date, shift_type):
    start, end = production_shift_bounds(selected_date, shift_type)
    return {
        'start': start,
        'end': end,
        'label': 'Первая смена' if shift_type == ShiftType.DAY else 'Вторая смена',
        'short_label': 'первая' if shift_type == ShiftType.DAY else 'вторая',
        'time_range': f'{start:%H:%M}–{end:%H:%M}',
    }


def shift_trip_filter(selected_date, shift_type):
    shift_start, shift_end = production_shift_bounds(selected_date, shift_type)
    assigned_start, assigned_end = assigned_shift_open_bounds(selected_date, shift_type)
    return (
        Q(
            loading_shift__shift_type=shift_type,
            loading_shift__opened_at__gte=assigned_start,
            loading_shift__opened_at__lt=assigned_end,
        )
        | Q(
            loading_shift__isnull=True,
            completed_at__gte=shift_start,
            completed_at__lt=shift_end,
        )
    )


def load_shift_trips(selected_date, shift_type, statuses=(TripStatus.COMPLETED,)):
    return list(
        Trip.objects
        .filter(status__in=statuses)
        .filter(shift_trip_filter(selected_date, shift_type))
        .select_related(
            'truck__equipment_type',
            'truck__model',
            'excavator__equipment_type',
            'excavator__model',
            'rock_type',
            'dump_point',
            'assigned_dump_point',
            'actual_dump_point',
            'loading_shift',
            'unloading_shift',
        )
        .order_by('completed_at', 'id')
    )


def load_shift_downtimes(selected_date, shift_type):
    shift_start, shift_end = production_shift_bounds(selected_date, shift_type)
    return list(
        DowntimeEvent.objects
        .filter(started_at__lt=shift_end)
        .filter(Q(ended_at__isnull=True) | Q(ended_at__gt=shift_start))
        .select_related('equipment__equipment_type', 'equipment__model', 'reason')
        .order_by('started_at', 'id')
    )


def downtime_overlap_seconds(event, shift_start, shift_end, *, now=None):
    now = now or timezone.now()
    event_end = event.ended_at or min(now, shift_end)
    effective_start = max(event.started_at, shift_start)
    effective_end = min(event_end, shift_end)
    return max(int((effective_end - effective_start).total_seconds()), 0)


def format_downtime_total(reason_name, seconds, *, is_open=False):
    minutes = max(seconds // 60, 0)
    hours, minute_remainder = divmod(minutes, 60)
    duration_parts = []
    if hours:
        duration_parts.append(f'{hours}ч')
    if minute_remainder or not duration_parts:
        duration_parts.append(f'{minute_remainder} мин')
    open_mark = ' (открыт)' if is_open else ''
    return f'{reason_name}—{" ".join(duration_parts)}{open_mark}'


def notes_by_equipment(downtimes, shift_start, shift_end):
    now = timezone.now()
    totals = {}
    order = []
    for event in downtimes:
        key = (event.equipment_id, event.reason_id)
        if key not in totals:
            totals[key] = {
                'equipment_id': event.equipment_id,
                'reason_name': event.reason.name,
                'seconds': 0,
                'is_open': False,
            }
            order.append(key)
        totals[key]['seconds'] += downtime_overlap_seconds(
            event,
            shift_start,
            shift_end,
            now=now,
        )
        totals[key]['is_open'] = totals[key]['is_open'] or event.ended_at is None

    result = defaultdict(list)
    for key in order:
        total = totals[key]
        if total['seconds'] < 60:
            continue
        result[total['equipment_id']].append(format_downtime_total(
            total['reason_name'],
            total['seconds'],
            is_open=total['is_open'],
        ))
    return result


def append_unique(target, value):
    value = (value or '').strip()
    if value and value not in target:
        target.append(value)


def effective_dump_point(trip):
    return trip.actual_dump_point or trip.assigned_dump_point or trip.dump_point


def build_truck_rows(trips, downtimes, meta):
    grouped = defaultdict(list)
    for trip in trips:
        grouped[trip.truck_id].append(trip)

    # SQLite does not case-fold Cyrillic for LIKE, so `icontains` would hide
    # "Самосвал" in local/test runs. Classify the small active roster
    # in Python to keep the same result on SQLite and PostgreSQL.
    active_trucks = [
        equipment
        for equipment in Equipment.objects.filter(is_active=True).select_related('equipment_type', 'model')
        if 'самосвал' in equipment.equipment_type.name.lower()
    ]
    known_ids = {equipment.id for equipment in active_trucks}
    active_trucks.extend(
        trip.truck for trip in trips
        if trip.truck_id not in known_ids and not known_ids.add(trip.truck_id)
    )
    notes_map = notes_by_equipment(downtimes, meta['start'], meta['end'])

    rows = []
    for truck in sorted(active_trucks, key=lambda item: (is_nhl(item), natural_key(item.garage_number))):
        truck_trips = grouped.get(truck.id, [])
        volume = decimal_sum(trip.volume_m3 for trip in truck_trips)
        m3km = decimal_sum(
            (trip.volume_m3 or ZERO) * (trip.transport_distance_km or ZERO)
            for trip in truck_trips
        )
        distance = (m3km / volume).quantize(ONE_DECIMAL, rounding=ROUND_HALF_UP) if volume else None
        notes = list(notes_map.get(truck.id, []))
        for trip in truck_trips:
            append_unique(notes, trip.downtime_text)
            append_unique(notes, trip.note)
        rows.append({
            'equipment': truck,
            'label': equipment_label(truck),
            'fleet': 'nhl' if is_nhl(truck) else 'belaz',
            'trip_count': len(truck_trips),
            'distance': distance,
            'volume': volume,
            'm3km': m3km,
            'notes': '; '.join(notes),
            'idle': not truck_trips,
        })

    totals = {
        'trip_count': sum(row['trip_count'] for row in rows),
        'volume': decimal_sum(row['volume'] for row in rows),
        'm3km': decimal_sum(row['m3km'] for row in rows),
    }
    totals['distance'] = (
        (totals['m3km'] / totals['volume']).quantize(ONE_DECIMAL, rounding=ROUND_HALF_UP)
        if totals['volume'] else None
    )
    return rows, totals


def build_excavation_rows(trips, downtimes, meta):
    grouped = {}
    for trip in trips:
        dump_point = effective_dump_point(trip)
        key = (
            trip.rock_type_id,
            trip.excavator_id,
            trip.planned_volume_m3,
            trip.loading_horizon,
            trip.loading_block,
            dump_point.id if dump_point else None,
            trip.transport_distance_km,
        )
        if key not in grouped:
            grouped[key] = {
                'rock_type': str(trip.rock_type),
                'excavator': trip.excavator,
                'planned_volume': trip.planned_volume_m3,
                'volume': ZERO,
                'horizon': trip.loading_horizon,
                'block': trip.loading_block,
                'dump_point': str(dump_point) if dump_point else '',
                'distance': trip.transport_distance_km,
                'trip_count': 0,
                'trip_notes': [],
            }
        row = grouped[key]
        row['volume'] += trip.volume_m3 or ZERO
        row['trip_count'] += 1
        append_unique(row['trip_notes'], trip.downtime_text)
        append_unique(row['trip_notes'], trip.note)

    notes_map = notes_by_equipment(downtimes, meta['start'], meta['end'])
    rows = []
    for row in grouped.values():
        notes = list(notes_map.get(row['excavator'].id, []))
        for value in row.pop('trip_notes'):
            append_unique(notes, value)
        row['excavator_label'] = row['excavator'].garage_number
        row['downtime'] = '; '.join(notes_map.get(row['excavator'].id, []))
        row['note'] = '; '.join(value for value in notes if value not in notes_map.get(row['excavator'].id, []))
        rows.append(row)

    rows.sort(key=lambda row: (row['rock_type'].lower(), natural_key(row['excavator_label']), row['dump_point'].lower()))
    downtime_shown_for = set()
    for row in rows:
        equipment_id = row['excavator'].id
        if equipment_id in downtime_shown_for:
            row['downtime'] = ''
            continue
        downtime_shown_for.add(equipment_id)
    totals = {
        'planned_volume': decimal_sum(row['planned_volume'] for row in rows),
        'volume': decimal_sum(row['volume'] for row in rows),
        'trip_count': sum(row['trip_count'] for row in rows),
    }

    trip_excavator_ids = {trip.excavator_id for trip in trips}
    other_equipment = {}
    for event in downtimes:
        if event.equipment_id not in trip_excavator_ids and 'самосвал' not in event.equipment.equipment_type.name.lower():
            other_equipment[event.equipment_id] = event.equipment
    shift_equipment = (
        EmployeeShift.objects
        .filter(shift_type=meta.get('shift_type'), opened_at__gte=meta['start'], opened_at__lt=meta['end'])
        .exclude(equipment__isnull=True)
        .select_related('equipment__equipment_type', 'equipment__model')
    )
    for employee_shift in shift_equipment:
        equipment = employee_shift.equipment
        if equipment.id not in trip_excavator_ids and 'самосвал' not in equipment.equipment_type.name.lower():
            other_equipment[equipment.id] = equipment

    other_rows = []
    for equipment in sorted(other_equipment.values(), key=lambda item: natural_key(item.garage_number)):
        notes = notes_map.get(equipment.id, [])
        other_rows.append({
            'equipment': equipment,
            'label': f'{equipment.model.name if equipment.model_id else equipment.equipment_type.name} {equipment.garage_number}'.strip(),
            'note': '; '.join(notes) or 'Работа в смене без завершённых рейсов',
        })
    return rows, totals, other_rows


def build_hourly_shift_matrix(trips, meta):
    """Build the dispatcher hourly matrix by loading time at the excavator."""
    hours = []
    hour_start = meta['start']
    while hour_start < meta['end']:
        hour_end = min(hour_start + timedelta(hours=1), meta['end'])
        hours.append({
            'index': len(hours),
            'start': hour_start,
            'end': hour_end,
            'label': timezone.localtime(hour_end).strftime('%H.%M'),
            'range_label': (
                f'{timezone.localtime(hour_start):%H:%M}–'
                f'{timezone.localtime(hour_end):%H:%M}'
            ),
            'points': [],
            'trip_count': 0,
            'volume': ZERO,
        })
        hour_start = hour_end

    bucketed = defaultdict(lambda: {'trip_count': 0, 'volume': ZERO})
    point_objects = {}
    excavator_objects = {}
    for trip in trips:
        loaded_at = trip.created_at
        hour_index = next((
            hour['index'] for hour in hours
            if hour['start'] <= loaded_at < hour['end']
        ), None)
        if hour_index is None:
            continue
        point = effective_dump_point(trip)
        point_key = point.id if point else 0
        point_objects[point_key] = point
        excavator_objects[trip.excavator_id] = trip.excavator
        fleet = 'nhl' if is_nhl(trip.truck) else 'belaz'
        metric = bucketed[(trip.excavator_id, fleet, hour_index, point_key)]
        metric['trip_count'] += 1
        metric['volume'] += trip.volume_m3 or ZERO
        hours[hour_index]['trip_count'] += 1
        hours[hour_index]['volume'] += trip.volume_m3 or ZERO

    points_by_hour = defaultdict(set)
    for _, _, hour_index, point_key in bucketed:
        points_by_hour[hour_index].add(point_key)
    for hour in hours:
        hour['points'] = [
            {
                'key': point_key,
                'label': str(point_objects[point_key]) if point_objects[point_key] else 'Без точки',
            }
            for point_key in sorted(
                points_by_hour[hour['index']],
                key=lambda key: natural_key(str(point_objects[key]) if point_objects[key] else 'Без точки'),
            )
        ]
        hour['point_totals'] = [
            {
                'trip_count': sum(
                    metric['trip_count']
                    for (excavator_id, fleet, hour_index, point_key), metric in bucketed.items()
                    if hour_index == hour['index'] and point_key == point['key']
                ),
                'volume': decimal_sum(
                    metric['volume']
                    for (excavator_id, fleet, hour_index, point_key), metric in bucketed.items()
                    if hour_index == hour['index'] and point_key == point['key']
                ),
            }
            for point in hour['points']
        ]

    active_excavators = [
        equipment
        for equipment in Equipment.objects.filter(is_active=True).select_related('equipment_type', 'model')
        if 'экскаватор' in equipment.equipment_type.name.lower()
    ]
    known_ids = {equipment.id for equipment in active_excavators}
    active_excavators.extend(
        equipment for equipment in excavator_objects.values()
        if equipment.id not in known_ids and not known_ids.add(equipment.id)
    )

    rows = []
    grand_trip_count = 0
    grand_volume = ZERO
    for excavator in sorted(active_excavators, key=lambda item: natural_key(item.garage_number)):
        fleet_rows = []
        excavator_hour_totals = []
        excavator_trip_count = 0
        excavator_volume = ZERO
        for fleet, fleet_label in (('belaz', 'БелАЗ'), ('nhl', 'NHL')):
            fleet_hours = []
            fleet_trip_count = 0
            fleet_volume = ZERO
            for hour in hours:
                cells = []
                hour_trip_count = 0
                hour_volume = ZERO
                for point in hour['points']:
                    metric = bucketed[(excavator.id, fleet, hour['index'], point['key'])]
                    cells.append(metric)
                    hour_trip_count += metric['trip_count']
                    hour_volume += metric['volume']
                fleet_hours.append({
                    'cells': cells,
                    'trip_count': hour_trip_count,
                    'volume': hour_volume,
                })
                fleet_trip_count += hour_trip_count
                fleet_volume += hour_volume
            fleet_rows.append({
                'key': fleet,
                'label': fleet_label,
                'hours': fleet_hours,
                'trip_count': fleet_trip_count,
                'volume': fleet_volume,
            })
            excavator_trip_count += fleet_trip_count
            excavator_volume += fleet_volume

        for hour_index, hour in enumerate(hours):
            point_totals = []
            for point_index, _point in enumerate(hour['points']):
                point_totals.append({
                    'trip_count': sum(row['hours'][hour_index]['cells'][point_index]['trip_count'] for row in fleet_rows),
                    'volume': decimal_sum(row['hours'][hour_index]['cells'][point_index]['volume'] for row in fleet_rows),
                })
            excavator_hour_totals.append({
                'cells': point_totals,
                'trip_count': sum(row['hours'][hour_index]['trip_count'] for row in fleet_rows),
                'volume': decimal_sum(row['hours'][hour_index]['volume'] for row in fleet_rows),
            })

        rows.append({
            'equipment': excavator,
            'label': excavator.garage_number,
            'fleet_rows': fleet_rows,
            'hours': excavator_hour_totals,
            'trip_count': excavator_trip_count,
            'volume': excavator_volume,
        })
        grand_trip_count += excavator_trip_count
        grand_volume += excavator_volume

    return {
        'hours': hours,
        'rows': rows,
        'trip_count': grand_trip_count,
        'volume': grand_volume,
        'column_count': 4 + sum(len(hour['points']) + 2 for hour in hours),
        'min_width_px': 300 + sum((len(hour['points']) * 74) + 126 for hour in hours),
    }


def build_quality_issues(trips, open_trips, downtimes):
    issues = []
    missing_volume = sum(1 for trip in trips if trip.volume_m3 is None)
    missing_distance = sum(1 for trip in trips if trip.transport_distance_km is None)
    open_downtimes = sum(1 for event in downtimes if event.ended_at is None)
    if open_trips:
        issues.append({'level': 'danger', 'label': 'Незакрытые рейсы', 'count': len(open_trips)})
    if missing_volume:
        issues.append({'level': 'danger', 'label': 'Рейсы без объёма', 'count': missing_volume})
    if missing_distance:
        issues.append({'level': 'risk', 'label': 'Рейсы без плеча', 'count': missing_distance})
    if open_downtimes:
        issues.append({'level': 'risk', 'label': 'Открытые простои', 'count': open_downtimes})
    return issues


def build_dispatcher_shift_report(selected_date, shift_type):
    meta = shift_meta(selected_date, shift_type)
    meta['shift_type'] = shift_type
    trips = load_shift_trips(selected_date, shift_type)
    hourly_trips = load_shift_trips(
        selected_date,
        shift_type,
        statuses=(TripStatus.LOADED_WAITING_UNLOAD, TripStatus.COMPLETED),
    )
    open_trips = load_shift_trips(selected_date, shift_type, statuses=OPEN_TRIP_STATUSES)
    downtimes = load_shift_downtimes(selected_date, shift_type)
    truck_rows, truck_totals = build_truck_rows(trips, downtimes, meta)
    excavation_rows, excavation_totals, other_rows = build_excavation_rows(trips, downtimes, meta)
    hourly = build_hourly_shift_matrix(hourly_trips, meta)
    return {
        'date': selected_date,
        'date_value': selected_date.isoformat(),
        'meta': meta,
        'trips': trips,
        'hourly_trips': hourly_trips,
        'open_trips': open_trips,
        'downtimes': downtimes,
        'truck_rows': truck_rows,
        'truck_totals': truck_totals,
        'excavation_rows': excavation_rows,
        'excavation_totals': excavation_totals,
        'other_rows': other_rows,
        'hourly': hourly,
        'issues': build_quality_issues(trips, open_trips, downtimes),
    }


def setup_print_sheet(sheet, widths):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A4'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def style_range(sheet, cell_range, fill=None, font=None, alignment=None, border=None):
    for row in sheet[cell_range]:
        for cell in row:
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if alignment:
                cell.alignment = alignment
            if border:
                cell.border = border


def write_truck_sheet(sheet, report):
    sheet.title = 'Самосвалы'
    setup_print_sheet(sheet, {'A': 14, 'B': 11, 'C': 11, 'D': 14, 'E': 14, 'F': 62})
    sheet.merge_cells('A1:F1')
    sheet['A1'] = f'Итоги смены по самосвалам за {report["date"]:%d.%m.%Y}'
    sheet.merge_cells('A2:F2')
    sheet['A2'] = f'{report["meta"]["label"]} смена ({report["meta"]["time_range"]})'
    headers = ['№ А/С', 'Рейсы', 'км', 'Итог, м³', 'м³×км', 'Примечание / простои']
    sheet.append(headers)
    yellow = PatternFill('solid', fgColor='FFC000')
    blue = PatternFill('solid', fgColor='D9E2F3')
    green = PatternFill('solid', fgColor='C6E0B4')
    total_fill = PatternFill('solid', fgColor='FFE699')
    border = Border(left=THIN_BLACK, right=THIN_BLACK, top=THIN_BLACK, bottom=THIN_BLACK)
    style_range(sheet, 'A1:F1', fill=yellow, font=Font(bold=True, size=14), alignment=Alignment(horizontal='center'))
    style_range(sheet, 'A2:F2', fill=yellow, font=Font(bold=True), alignment=Alignment(horizontal='center'))
    style_range(sheet, 'A3:F3', fill=yellow, font=Font(bold=True), alignment=Alignment(horizontal='center', wrap_text=True), border=border)
    sheet.row_dimensions[3].height = 32

    for row_data in report['truck_rows']:
        sheet.append([
            row_data['label'],
            row_data['trip_count'] or '',
            float(row_data['distance']) if row_data['distance'] is not None else '',
            float(row_data['volume']) if row_data['volume'] else '',
            float(row_data['m3km']) if row_data['m3km'] else '',
            row_data['notes'],
        ])
        row_number = sheet.max_row
        fill = green if row_data['fleet'] == 'nhl' else blue
        for cell in sheet[row_number]:
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.cell(row_number, 1).font = Font(bold=True, color='FF0000' if row_data['idle'] and row_data['notes'] else '000000')
        if row_data['idle'] and row_data['notes']:
            sheet.cell(row_number, 6).font = Font(bold=True, color='FF0000')
        sheet.cell(row_number, 3).number_format = '0.0'
        sheet.cell(row_number, 4).number_format = '#,##0'
        sheet.cell(row_number, 5).number_format = '#,##0'

    total = report['truck_totals']
    sheet.append([
        'Итого',
        total['trip_count'],
        float(total['distance']) if total['distance'] is not None else '',
        float(total['volume']),
        float(total['m3km']),
        '',
    ])
    total_row = sheet.max_row
    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.border = Border(left=MEDIUM_BLACK, right=MEDIUM_BLACK, top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.cell(total_row, 3).number_format = '0.0'
    sheet.cell(total_row, 4).number_format = '#,##0'
    sheet.cell(total_row, 5).number_format = '#,##0'
    sheet.print_area = f'A1:F{total_row}'


def write_excavation_sheet(sheet, report):
    sheet.title = 'Выемочное оборудование'
    setup_print_sheet(sheet, {'A': 23, 'B': 18, 'C': 15, 'D': 17, 'E': 12, 'F': 12, 'G': 25, 'H': 15, 'I': 25, 'J': 35})
    sheet.merge_cells('A1:J1')
    sheet['A1'] = f'Отчёт о работе выемочного оборудования за {report["date"]:%d.%m.%Y}'
    sheet.merge_cells('A2:J2')
    sheet['A2'] = f'{report["meta"]["label"]} смена ({report["meta"]["time_range"]})'
    headers = ['Тип грунта', 'Наименование / номер экскаватора', 'План, м³', 'Выполнено, м³', 'Горизонт', 'Блок', 'Место разгрузки', 'Плечо, км', 'Простои', 'Примечание']
    sheet.append(headers)
    blue = PatternFill('solid', fgColor='D9E2F3')
    other_fill = PatternFill('solid', fgColor='BDD7EE')
    total_fill = PatternFill('solid', fgColor='FFE699')
    border = Border(left=THIN_BLACK, right=THIN_BLACK, top=THIN_BLACK, bottom=THIN_BLACK)
    style_range(sheet, 'A1:J1', fill=PatternFill('solid', fgColor='9DC3E6'), font=Font(bold=True, size=13), alignment=Alignment(horizontal='center'))
    style_range(sheet, 'A2:J2', fill=blue, font=Font(bold=True), alignment=Alignment(horizontal='center'))
    style_range(sheet, 'A3:J3', font=Font(bold=True), alignment=Alignment(horizontal='center', vertical='center', wrap_text=True), border=border)
    sheet.row_dimensions[3].height = 56

    for row_data in report['excavation_rows']:
        sheet.append([
            row_data['rock_type'], row_data['excavator_label'],
            float(row_data['planned_volume']) if row_data['planned_volume'] is not None else '',
            float(row_data['volume']), row_data['horizon'], row_data['block'], row_data['dump_point'],
            float(row_data['distance']) if row_data['distance'] is not None else '',
            row_data['downtime'], row_data['note'],
        ])
        row_number = sheet.max_row
        for cell in sheet[row_number]:
            cell.fill = blue
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.cell(row_number, 3).number_format = '#,##0'
        sheet.cell(row_number, 4).number_format = '#,##0'
        sheet.cell(row_number, 8).number_format = '0.0'

    for other_row in report['other_rows']:
        sheet.append(['Прочие работы', other_row['label'], '', '', '', '', '', '', '', other_row['note']])
        row_number = sheet.max_row
        for cell in sheet[row_number]:
            cell.fill = other_fill
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    totals = report['excavation_totals']
    sheet.append(['Итого', '', float(totals['planned_volume']), float(totals['volume']), '', '', '', '', '', f'{totals["trip_count"]} рейсов'])
    total_row = sheet.max_row
    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.border = Border(left=MEDIUM_BLACK, right=MEDIUM_BLACK, top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.print_area = f'A1:J{total_row}'


def write_hourly_sheet(sheet, report):
    sheet.title = 'Почасовая сводка'
    hourly = report['hourly']
    hours = hourly['hours']
    border = Border(left=THIN_BLACK, right=THIN_BLACK, top=THIN_BLACK, bottom=THIN_BLACK)
    medium_border = Border(left=MEDIUM_BLACK, right=MEDIUM_BLACK, top=MEDIUM_BLACK, bottom=MEDIUM_BLACK)
    belaz_fill = PatternFill('solid', fgColor='E7E6E6')
    nhl_fill = PatternFill('solid', fgColor='BDD7EE')
    total_fill = PatternFill('solid', fgColor='FFD966')
    hour_fill = PatternFill('solid', fgColor='FFF2CC')

    columns = []
    current_column = 3
    for hour in hours:
        point_columns = []
        for point in hour['points']:
            point_columns.append({'column': current_column, 'point': point})
            current_column += 1
        columns.append({
            'hour': hour,
            'points': point_columns,
            'trip_column': current_column,
            'volume_column': current_column + 1,
            'start_column': point_columns[0]['column'] if point_columns else current_column,
            'end_column': current_column + 1,
        })
        current_column += 2
    total_trip_column = current_column
    total_volume_column = current_column + 1
    last_column = total_volume_column

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.cell(1, 1, 'Почасовая сводка смены')
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(2, 1, f'{report["date"]:%d.%m.%Y} · {report["meta"]["label"]} смена {report["meta"]["time_range"]}')
    sheet.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
    sheet.merge_cells(start_row=3, start_column=2, end_row=4, end_column=2)
    sheet.cell(3, 1, '№ экс')
    sheet.cell(3, 2, 'а/с')
    for column in columns:
        sheet.merge_cells(
            start_row=3,
            start_column=column['start_column'],
            end_row=3,
            end_column=column['end_column'],
        )
        sheet.cell(3, column['start_column'], column['hour']['label'])
        for point_column in column['points']:
            sheet.cell(4, point_column['column'], point_column['point']['label'])
        sheet.cell(4, column['trip_column'], 'Рейсы')
        sheet.cell(4, column['volume_column'], 'Объём')
    sheet.merge_cells(start_row=3, start_column=total_trip_column, end_row=3, end_column=total_volume_column)
    sheet.cell(3, total_trip_column, 'ИТОГО')
    sheet.cell(4, total_trip_column, 'Рейсы')
    sheet.cell(4, total_volume_column, 'Объём')

    style_range(sheet, f'A1:{get_column_letter(last_column)}1', fill=total_fill, font=Font(bold=True, size=15), alignment=Alignment(horizontal='center'))
    style_range(sheet, f'A2:{get_column_letter(last_column)}2', fill=total_fill, font=Font(bold=True), alignment=Alignment(horizontal='center'))
    style_range(sheet, f'A3:{get_column_letter(last_column)}4', fill=hour_fill, font=Font(bold=True), alignment=Alignment(horizontal='center', vertical='center', wrap_text=True), border=border)
    sheet.freeze_panes = 'C5'
    sheet.sheet_view.showGridLines = False

    detail_rows = []
    for row in hourly['rows']:
        first_row = sheet.max_row + 1
        fleet_excel_rows = []
        for fleet_row in row['fleet_rows']:
            sheet.append(['', fleet_row['label']])
            row_number = sheet.max_row
            fleet_excel_rows.append(row_number)
            for hour_index, column in enumerate(columns):
                hour_data = fleet_row['hours'][hour_index]
                for point_index, point_column in enumerate(column['points']):
                    value = hour_data['cells'][point_index]['trip_count']
                    sheet.cell(row_number, point_column['column'], value or '')
                sheet.cell(row_number, column['trip_column'], hour_data['trip_count'] or '')
                sheet.cell(row_number, column['volume_column'], float(hour_data['volume']) if hour_data['volume'] else '')
            sheet.cell(row_number, total_trip_column, fleet_row['trip_count'] or '')
            sheet.cell(row_number, total_volume_column, float(fleet_row['volume']) if fleet_row['volume'] else '')
            row_fill = nhl_fill if fleet_row['key'] == 'nhl' else belaz_fill
            style_range(sheet, f'A{row_number}:{get_column_letter(last_column)}{row_number}', fill=row_fill, alignment=Alignment(horizontal='center', vertical='center'), border=border)
            detail_rows.append(row_number)
        sheet.merge_cells(start_row=first_row, start_column=1, end_row=first_row + 1, end_column=1)
        sheet.cell(first_row, 1, row['label'])
        sheet.cell(first_row, 1).font = Font(bold=True, size=12)
        sheet.cell(first_row, 1).alignment = Alignment(horizontal='center', vertical='center')

        sheet.append(['Рейсы', 'Объёмы'])
        subtotal_row = sheet.max_row
        for column in columns:
            for point_column in column['points']:
                letter = get_column_letter(point_column['column'])
                sheet.cell(subtotal_row, point_column['column'], f'=SUM({letter}{fleet_excel_rows[0]}:{letter}{fleet_excel_rows[-1]})')
            trip_letter = get_column_letter(column['trip_column'])
            volume_letter = get_column_letter(column['volume_column'])
            sheet.cell(subtotal_row, column['trip_column'], f'=SUM({trip_letter}{fleet_excel_rows[0]}:{trip_letter}{fleet_excel_rows[-1]})')
            sheet.cell(subtotal_row, column['volume_column'], f'=SUM({volume_letter}{fleet_excel_rows[0]}:{volume_letter}{fleet_excel_rows[-1]})')
        total_trip_letter = get_column_letter(total_trip_column)
        total_volume_letter = get_column_letter(total_volume_column)
        sheet.cell(subtotal_row, total_trip_column, f'=SUM({total_trip_letter}{fleet_excel_rows[0]}:{total_trip_letter}{fleet_excel_rows[-1]})')
        sheet.cell(subtotal_row, total_volume_column, f'=SUM({total_volume_letter}{fleet_excel_rows[0]}:{total_volume_letter}{fleet_excel_rows[-1]})')
        style_range(sheet, f'A{subtotal_row}:{get_column_letter(last_column)}{subtotal_row}', fill=total_fill, font=Font(bold=True), alignment=Alignment(horizontal='center', vertical='center'), border=medium_border)

    sheet.append(['ИТОГО ЗА СМЕНУ', ''])
    grand_row = sheet.max_row
    for column in columns:
        for point_column in column['points']:
            letter = get_column_letter(point_column['column'])
            sheet.cell(grand_row, point_column['column'], f'=SUM({",".join(f"{letter}{row}" for row in detail_rows)})' if detail_rows else 0)
        sheet.cell(grand_row, column['trip_column'], column['hour']['trip_count'])
        sheet.cell(grand_row, column['volume_column'], float(column['hour']['volume']))
    sheet.cell(grand_row, total_trip_column, hourly['trip_count'])
    sheet.cell(grand_row, total_volume_column, float(hourly['volume']))
    style_range(sheet, f'A{grand_row}:{get_column_letter(last_column)}{grand_row}', fill=total_fill, font=Font(bold=True, size=11), alignment=Alignment(horizontal='center', vertical='center'), border=medium_border)

    sheet.column_dimensions['A'].width = 11
    sheet.column_dimensions['B'].width = 11
    for column in range(3, last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 11
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '1:4'
    sheet.print_area = f'A1:{get_column_letter(last_column)}{grand_row}'


def build_shift_report_workbook(report, report_kind):
    workbook = Workbook()
    if report_kind == 'hourly':
        write_hourly_sheet(workbook.active, report)
    elif report_kind == 'excavation':
        write_excavation_sheet(workbook.active, report)
    elif report_kind == 'all':
        write_truck_sheet(workbook.active, report)
        write_excavation_sheet(workbook.create_sheet(), report)
    else:
        write_truck_sheet(workbook.active, report)
    return workbook
